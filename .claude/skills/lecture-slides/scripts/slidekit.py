"""강의 슬라이드 도우미 (python-pptx + 프로젝트 템플릿).

사용 예 (lectures/N강/build_slides.py):

    import sys, pathlib
    ROOT = pathlib.Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(ROOT / ".claude/skills/lecture-slides/scripts"))
    from slidekit import *

    deck = Deck(__file__, "1강_확률변수와_확률분포.pptx",
                footer=("확률과 통계", "Ⅲ. 통계 › 1. 확률분포"))
    deck.cover("확률변수와 확률분포", "1강 · Ⅲ. 통계 › 1. 확률분포", notes="...")
    deck.ending("수고하셨습니다", "다음 시간: ...")
    s = deck.slide("학습 목표", "오늘 수업을 마치면 ", "이렇게 할 수 있어요", notes="...")
    text(s, L, 1650000, CW, 500000, "본문 [[강조]] x_{i}", 18)
    deck.credits_slide()
    deck.save()          # 저장 + 영상 자동 재생 설정

좌표 단위는 EMU (1인치 = 914400). 슬라이드 12192000 x 6858000.
본문 영역: x = L..R, y ≈ 1550000..6200000 (제목 위, 바닥글 아래 6380000).
텍스트 마크업: **굵게**, [[강조색]], x_{i} 아래첨자.
"""
import json
import os
import re
import subprocess

from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: F401  (build 스크립트에서 사용)
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SKILL_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.abspath(os.path.join(SKILL_DIR, "..", "..", "..", ".."))
TEMPLATE = os.path.join(PROJECT, "references", "제이비엠스토리 피피티 무료템플릿_251111.pptx")
TPL_COVER, TPL_END = 0, 22          # 템플릿에서 남길 표지 / 마무리 슬라이드 (0부터)

BG = RGBColor(0x22, 0x2A, 0x35)
CARD = RGBColor(0x2C, 0x37, 0x46)
CARD2 = RGBColor(0x34, 0x41, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB4, 0xC0, 0xCE)
A1, A2, A3, A4, A5 = (RGBColor.from_string(c) for c in ("1F608B", "2980B9", "4098D4", "7BB8E1", "9FCBE9"))
FONT = "Pretendard"

SW, SH = 12192000, 6858000
IN = 914400
L, R = 838200, 11353800             # 본문 좌우 경계
CW = R - L

_ctx = {"img": None, "anim": None, "credits": {}}


# ---------------------------------------------------------------- 텍스트
def set_font(run, size, color=WHITE, bold=False, light=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    name = FONT + (" Light" if light else "")
    f.name = name
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "ko-KR")        # 한글 단어 단위 줄바꿈
    rPr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def add_runs(p, text, size, color=WHITE, bold=False, light=False):
    """x_{i} 아래첨자, **굵게**, [[강조색]] 지원."""
    for tok in re.split(r"(_\{[^}]*\}|\*\*[^*]+\*\*|\[\[[^\]]+\]\])", text):
        if not tok:
            continue
        r = p.add_run()
        if tok.startswith("_{"):
            r.text = tok[2:-1]
            set_font(r, size, color, bold, light)
            r._r.get_or_add_rPr().set("baseline", "-25000")
        elif tok.startswith("**"):
            r.text = tok[2:-2]
            set_font(r, size, color, True, False)
        elif tok.startswith("[["):
            r.text = tok[2:-2]
            set_font(r, size, A4, True, False)
        else:
            r.text = tok
            set_font(r, size, color, bold, light)


def text(slide, x, y, w, h, lines, size=18, color=WHITE, align=PP_ALIGN.LEFT, bold=False,
         light=False, anchor=MSO_ANCHOR.TOP, spacing=1.2, after=6):
    """lines: 문자열 또는 문자열 리스트(문단)."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(after)
        p._p.get_or_add_pPr().set("latinLnBrk", "0")
        add_runs(p, line, size, color, bold, light)
    return tb


# ---------------------------------------------------------------- 도형
def box(slide, x, y, w, h, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
    s = slide.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def circle_num(slide, x, y, d, label, fill=A2, size=16):
    c = box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    tf = c.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, label, size, WHITE, True)
    return c


def table(slide, x, y, w, rows, col0w=None, rowh=int(0.55 * IN), size=18):
    """rows[0]은 머리행, 첫 열은 강조색."""
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Emu(x), Emu(y), Emu(w), Emu(rowh * nr)).table
    col0w = col0w or w // nc
    gt.columns[0].width = Emu(col0w)
    for j in range(1, nc):
        gt.columns[j].width = Emu((w - col0w) // (nc - 1))
    for i, row in enumerate(rows):
        gt.rows[i].height = Emu(rowh)
        for j, val in enumerate(row):
            cell = gt.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = A1 if j == 0 else (CARD if i % 2 == 0 else CARD2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = ""
            p.alignment = PP_ALIGN.CENTER
            add_runs(p, val, size, WHITE, j == 0)
    tblPr = gt._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    return gt


def numbered_cards(slide, items, x=L, y=1650000, w=CW, h=1000000, gap=150000, title_size=21, desc_size=15):
    """items: [(제목, 설명), ...] 번호 원 + 카드 목록."""
    colors = [A2, A3, A4, A5]
    for i, (t, d) in enumerate(items):
        yy = y + i * (h + gap)
        box(slide, x, yy, w, h)
        d_ = min(560000, h - 300000)
        circle_num(slide, x + 230000, yy + (h - d_) // 2, d_, str(i + 1), colors[i % 4], 20)
        tx = x + 230000 + d_ + 300000
        if d:
            text(slide, tx, yy + h // 2 - 430000, x + w - tx - 200000, 460000, t, title_size, bold=True,
                 anchor=MSO_ANCHOR.BOTTOM)
            text(slide, tx, yy + h // 2 + 60000, x + w - tx - 200000, 400000, d, desc_size, MUTED)
        else:
            text(slide, tx, yy, x + w - tx - 200000, h, t, title_size - 2, anchor=MSO_ANCHOR.MIDDLE)


def analogy_tag(slide, x, y, label="비유"):
    box(slide, x, y, 1150000, 340000, A2, radius=0.5)
    text(slide, x, y, 1150000, 340000, label, 14, WHITE, PP_ALIGN.CENTER, bold=True, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- 이미지 · 영상
def picture(slide, name, x, y, w, h, focus=(0.5, 0.5), credit=True):
    """images/<name>.jpg 를 영역에 꽉 차게 잘라 넣고 아래에 출처를 단다.
    focus: 잘라낼 때 남길 위치 (0=왼쪽/위, 1=오른쪽/아래)."""
    path = os.path.join(_ctx["img"], name + ".jpg")
    iw, ih = PILImage.open(path).size
    pic = slide.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    ta, ia = w / h, iw / ih
    if ia > ta:
        extra = 1 - ta / ia
        pic.crop_left = extra * focus[0]
        pic.crop_right = extra - pic.crop_left
    else:
        extra = 1 - ia / ta
        pic.crop_top = extra * focus[1]
        pic.crop_bottom = extra - pic.crop_top
    if credit:
        c = _ctx["credits"].get(name)
        if c:
            text(slide, x, y + h + 60000, w, 250000, f"사진: {c['creator_short']} / {c['license']}", 10, MUTED,
                 PP_ALIGN.RIGHT)
    return pic


def video(slide, name, x, y, w):
    """animations/<name>.mp4 (16:9)를 넣는다. 포스터는 <name>.png (마지막 장면). 높이 = w*9/16."""
    h = int(w * 9 / 16)
    return slide.shapes.add_movie(os.path.join(_ctx["anim"], name + ".mp4"), Emu(x), Emu(y), Emu(w), Emu(h),
                                  poster_frame_image=os.path.join(_ctx["anim"], name + ".png"),
                                  mime_type="video/mp4")


# ---------------------------------------------------------------- 덱
def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _replace_text(sh, new, size=None):
    p = sh.text_frame.paragraphs[0]
    r0 = p.runs[0]
    r0.text = new
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in sh.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if size:
        r0.font.size = Pt(size)
    r0.font.name = FONT
    rPr = r0._r.get_or_add_rPr()
    el = rPr.find(qn("a:ea"))
    if el is None:
        el = rPr.makeelement(qn("a:ea"), {})
        rPr.append(el)
    el.set("typeface", FONT)


class Deck:
    def __init__(self, build_file, out_name, footer=("", ""), template=TEMPLATE):
        self.dir = os.path.dirname(os.path.abspath(build_file))
        self.out = os.path.join(self.dir, out_name)
        self.footer = footer
        _ctx["img"] = os.path.join(self.dir, "images")
        _ctx["anim"] = os.path.join(self.dir, "animations")
        cpath = os.path.join(_ctx["img"], "credits.json")
        _ctx["credits"] = json.load(open(cpath, encoding="utf-8")) if os.path.exists(cpath) else {}

        self.prs = Presentation(template)
        lst = self.prs.slides._sldIdLst
        for i, sid in enumerate(list(lst)):
            if i not in (TPL_COVER, TPL_END):
                self.prs.part.drop_rel(sid.rId)
                lst.remove(sid)
        self._cover, self._end = self.prs.slides[0], self.prs.slides[1]

    def cover(self, title, subtitle, notes=""):
        for sh in _walk(self._cover.shapes):
            if sh.has_text_frame:
                t = sh.text_frame.text
                if t.strip() == "Introduction":
                    _replace_text(sh, title, 24 if len(title) <= 10 else 20)
                elif "Subtitle" in t:
                    _replace_text(sh, subtitle, 12)
        if notes:
            self._cover.notes_slide.notes_text_frame.text = notes

    def ending(self, title="수고하셨습니다", subtitle=""):
        for sh in _walk(self._end.shapes):
            if sh.has_text_frame:
                t = sh.text_frame.text
                if "THANK" in t:
                    _replace_text(sh, title, 40)
                elif "Attention" in t:
                    _replace_text(sh, subtitle, 16)

    def slide(self, title, subtitle_plain="", subtitle_accent="", notes=""):
        """제목 + 부제(일반/강조) + 바닥글 + 발표자 노트. 쪽 번호 원은 템플릿 레이아웃이 자동으로 그린다."""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        text(s, 0, 330000, SW, 700000, title, 36, WHITE, PP_ALIGN.CENTER, light=True)
        if subtitle_plain or subtitle_accent:
            tb = text(s, 0, 1010000, SW, 340000, "", 15, MUTED, PP_ALIGN.CENTER)
            p = tb.text_frame.paragraphs[0]
            add_runs(p, subtitle_plain, 15, MUTED)
            add_runs(p, subtitle_accent, 15, A3)
        if any(self.footer):
            tb = text(s, 541215, 6382260, 6000000, 300000, "", 12)
            p = tb.text_frame.paragraphs[0]
            add_runs(p, self.footer[0], 12, A3, True)
            if self.footer[1]:
                add_runs(p, "  ·  " + self.footer[1], 12, MUTED)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def credits_slide(self):
        if not _ctx["credits"]:
            return
        s = self.slide("이미지 출처", "모든 사진은 ", "CC 라이선스 / 퍼블릭 도메인",
                       "CC BY / BY-SA 사진은 출처를 표시하면 수업 자료로 사용할 수 있습니다.")
        lines = [f"• {k}.jpg — {v['title']} / {v['creator_short']} / {v['license']}"
                 for k, v in _ctx["credits"].items()]
        text(s, L, 1650000, CW, 3500000, lines, 14 if len(lines) <= 9 else 12, WHITE, after=8)
        text(s, L, 5300000, CW, 700000,
             [f"원본 링크: {os.path.relpath(os.path.join(_ctx['img'], '출처.md'), PROJECT)}",
              "애니메이션: manim (Community Edition)으로 제작"], 13, MUTED, after=4)

    def save(self, autoplay=True):
        lst = self.prs.slides._sldIdLst
        end_id = list(lst)[1]          # 마무리 슬라이드를 맨 뒤로
        lst.remove(end_id)
        lst.append(end_id)
        self.prs.save(self.out)
        print("saved", self.out, len(self.prs.slides), "slides")
        if autoplay and any(sh.shape_type == 16 for s in self.prs.slides for sh in s.shapes):
            r = subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                                os.path.join(SKILL_DIR, "autoplay.ps1"), self.out],
                               capture_output=True, text=True, encoding="utf-8", errors="replace")
            print((r.stdout or "").strip() or r.stderr.strip())
